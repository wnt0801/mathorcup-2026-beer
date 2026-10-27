"""Evidence-only audits: no new predictive model selection and no fabricated labels."""
from pathlib import Path
import json
import hashlib
import numpy as np
import pandas as pd
from scipy.spatial.distance import cdist,pdist
from openpyxl import load_workbook
from review_config import ROOT,OUT,FROZEN


def write(df,name):df.to_csv(OUT/name,index=False,encoding='utf-8-sig')


def coverage_diagnostics(real,fake,style_train,imputer,scaler):
    """All distances use train-fitted global scaling; all variance estimates use
    observed native values (ddof=1). Undefined small-sample statistics remain NA.
    """
    a=scaler.transform(imputer.transform(style_train))
    b=scaler.transform(imputer.transform(fake))
    h=scaler.transform(imputer.transform(real))
    p=a.shape[1]
    nearest=cdist(b,a)/np.sqrt(p)
    hold=cdist(h,a)/np.sqrt(p)
    selfdist=cdist(a,a)/np.sqrt(p);np.fill_diagonal(selfdist,np.inf)
    loo=selfdist.min(axis=1)
    radius=float(np.quantile(loo,.95)) if len(a)>1 else np.nan
    htogen=(cdist(h,b)/np.sqrt(p)).min(axis=1)
    vt=style_train.var(ddof=1);vf=fake.var(ddof=1);vh=real.var(ddof=1)
    eligible=vt.gt(1e-12)&vf.notna()
    ratio=vf[eligible]/vt[eligible]
    scale2=pd.Series(scaler.scale_**2,index=style_train.columns)
    denominator=float((vt[eligible]/scale2[eligible]).sum())
    aggregate=float((vf[eligible]/scale2[eligible]).sum()/denominator) if denominator>0 else np.nan
    return dict(
        n_generated=len(fake),variance_features=int(eligible.sum()),
        generated_train_variance_ratio_median=float(ratio.median()),
        generated_train_total_scaled_variance_ratio=aggregate,
        train_total_scaled_variance=float((vt/scale2).sum(min_count=1)),
        generated_total_scaled_variance=float((vf/scale2).sum(min_count=1)),
        holdout_total_scaled_variance=float((vh/scale2).sum(min_count=1)) if len(real)>=2 else np.nan,
        holdout_variance_status='estimable_but_small_sample' if len(real)>=2 else 'undefined_n_less_than_2',
        generated_pairwise_RMS_mean=float(pdist(b).mean()/np.sqrt(p)),
        training_pairwise_RMS_mean=float(pdist(a).mean()/np.sqrt(p)) if len(a)>1 else np.nan,
        generated_to_same_style_train_mean=float(nearest.min(1).mean()),
        generated_to_same_style_train_q95=float(np.quantile(nearest.min(1),.95)),
        holdout_to_same_style_train_mean=float(hold.min(1).mean()),
        training_leave_one_out_mean=float(loo.mean()) if len(a)>1 else np.nan,
        train_nn_radius95=radius,
        holdout_to_generated_mean=float(htogen.mean()),
        holdout_covered_at_train_nn95=float((htogen<=radius).mean()) if np.isfinite(radius) else np.nan,
        generated_in_train_nn95_fraction=float((nearest.min(1)<=radius).mean()) if np.isfinite(radius) else np.nan,
    )


def scoring_sources():
    a4=pd.read_csv(ROOT/'data/4.csv');a7=pd.read_csv(ROOT/'data/7.csv')
    targets=[c for c in a4 if c not in ['beer','beer_id','tasting_category_fine']]
    summary=pd.read_excel(ROOT/'data/3.xlsx',sheet_name='summary')
    wb=load_workbook(ROOT/'data/3.xlsx',data_only=False)
    summary_numeric=summary.drop(columns=[c for c in summary if str(c).lower() in ['beer','comments']]).select_dtypes('number')
    cards=[]
    for ws in wb:
        if not ws.title.startswith('beer_'):continue
        formulas=[{'cell':c.coordinate,'formula':c.value} for row in ws for c in row if c.data_type=='f']
        cards.append({'sheet':ws.title,'formula_count':len(formulas),'formula_examples':json.dumps(formulas[:4],ensure_ascii=False),'interpretation':'rubric/template, not verified completed panel labels'})
    template_differences={}
    base=wb['beer_1']
    for ws in wb:
        if ws.title.startswith('beer_'):
            template_differences[ws.title]=[{'cell':c.coordinate,'base':base[c.coordinate].value,'value':c.value} for row in ws for c in row if c.value!=base[c.coordinate].value]
    wb.close();write(pd.DataFrame(cards),'q1_scorecard_inspection.csv')
    (OUT/'q1_scorecard_template_differences.json').write_text(json.dumps(template_differences,ensure_ascii=False,indent=2),encoding='utf-8')
    workbook_metadata=[]
    for filename in [ROOT/'data/3.xlsx',ROOT/'data/8.xlsx',FROZEN/'reference_supplement.xlsx']:
        book=load_workbook(filename,data_only=False)
        workbook_metadata.append({'source':filename.relative_to(ROOT).as_posix(),'defined_names':list(book.defined_names),'hidden_sheets':[s.title for s in book if s.sheet_state!='visible']})
        book.close()
    (OUT/'q1_workbook_metadata.json').write_text(json.dumps(workbook_metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    write(a4[targets].describe().T.reset_index(names='target'),'q1_target_scale_summary.csv')
    tables=[]
    for filename in [ROOT/'data/8.xlsx',FROZEN/'reference_supplement.xlsx']:
        book=pd.ExcelFile(filename)
        for sheet in book.sheet_names:
            d=pd.read_excel(book,sheet_name=sheet)
            matched=[str(c) for c in d.columns if c in targets]
            tables.append({'source':filename.relative_to(ROOT).as_posix(),'sheet':sheet,'rows':len(d),'columns':len(d.columns),'matched_a4_target_columns':len(matched),'column_names':json.dumps([str(c) for c in d.columns],ensure_ascii=False)})
    write(pd.DataFrame(tables),'q1_source_sheet_inventory.csv')
    ref=pd.read_excel(FROZEN/'reference_supplement.xlsx',sheet_name='Supplementary File S4').set_index('beer')
    delta=(a4.set_index('beer')[targets]-ref[targets]).abs()
    result={
        'attachment3_card_sheets':len(cards),'attachment3_summary_rows':len(summary),
        'attachment3_summary_numeric_nonzero':int((summary_numeric.notna()&summary_numeric.ne(0)).to_numpy().sum()),
        'attachment3_summary_numeric_missing':int(summary_numeric.isna().to_numpy().sum()),
        'attachment3_cards_only_differ_in_beer_number':all(all(r['cell']=='S3' for r in rows) for rows in template_differences.values()),
        'attachment4_targets':len(targets),'attachment4_negative_cells':int(a4[targets].lt(0).to_numpy().sum()),
        'attachment4_noninteger_cells':int((~np.isclose(a4[targets],np.round(a4[targets]),atol=1e-8,rtol=0)).sum()),
        'attachment7_rows':len(a7),'attachment7_attribute_columns':len(a7.columns)-3,
        'attachment7_negative_cells':int(a7.select_dtypes('number').drop(columns=['beer','beer_id'],errors='ignore').lt(0).to_numpy().sum()),
        'public_s4_matches_a4_max_absolute_error':float(delta.max().max()),
        'public_supplement_md5':hashlib.md5((FROZEN/'reference_supplement.xlsx').read_bytes()).hexdigest(),
        'reliable_original_0to6_calibration_found':False,
        'calibration_status':'not_identifiable_from_supplied_materials',
        'output_definitions':{'predicted_standardized':'continuous target scale supplied in attachment4','predicted_relative_grade':'training/fold quantile rank category; NOT original intensity','actual_relative_grade':'observed standardized target categorized with the same training/fold cut points; NOT original label'},
        'evidence_limit':'Worksheet/header inspection and scorecard formulas do not provide a verified paired original panel-score table or its inverse-transform parameters. Validation-experiment results cannot be silently treated as original250 training labels.'
    }
    (OUT/'q1_score_scale_evidence.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    note=f'''# 第一问评分口径说明

## 已确认事实

- 附件3实际有{len(cards)}张beer评分卡及summary，summary有{len(summary)}行，感官数值非零单元格为{result['attachment3_summary_numeric_nonzero']}，缺失单元格为{result['attachment3_summary_numeric_missing']}。逐格比较12张卡，仅S3中的Beer#编号不同；0–6刻度和公式模板不能充当已完成的人工评分。
- 附件4有50个连续感官目标，其中{result['attachment4_negative_cells']}个负值、{result['attachment4_noninteger_cells']}个非整数值。题面明确其经过标准化/中心化。
- A/F命名项目实际为30维，外加20项其他感官项目合计50维，与题面“约40维”不完全一致。
- 附件7有{len(a7)}款、{len(a7.columns)-3}个非标识属性，含负值与缺失；不是与评分卡配对的原始0–6面板标签。
- 附件8逐sheet检查及已缓存的公开原研究29个sheet已重新核验；公开S4与附件4最大差{result['public_s4_matches_a4_max_absolute_error']:.3g}。没有找到可验证的250款配对原始评分与可逆标准化参数。公开添加实验的评分不等于原250款校准数据。

## 当前保留的输出

- `predicted_standardized`：附件4连续目标尺度预测。
- `predicted_relative_grade`：训练分位数相对0–6级。训练OOF使用各折拟合集的六个分位点；测试集使用200款训练样本的分位点。
- `actual_relative_grade`：观测连续值按相同训练界点分档，仅用于相对等级误差计算；它也不是原始评分。
- 未建立raw/original_score列，未进行min-max、平移、四舍五入后冒称还原。即使假设z=(r-mu)/sigma，未知mu与sigma也不能唯一反推r；按评分者标准化、聚合或类别化还可能不可逆。

## 缺失信息与需要主办方确认的问题

1. 第一问要求的是原始面板强度等级，还是允许明确标注的相对等级？如要求原始强度，请提供配对原始标签或官方校准表。
2. 标准化在哪一层级进行（每评分者、每属性、每场次或全体）？中心/尺度参数、公式、面板聚合顺序、缺失及离散项目的处理是什么？
3. 目标维度以实际Aroma15+Flavour15为核心，还是以summary全部50项为准？“约40维”有哪些明确项目？
4. 对原始分级如何计分：整数、面板平均强度、等级界点、边界同分与评估指标分别是什么？

当前结论是“现有资料不足以唯一恢复”，不是“评分口径已经解决”。正式审核可审阅标准化预测及相对等级，但不能把本版本当作已证明满足原始0–6评分要求。
'''
    (OUT/'Q1_SCORING_SCOPE.md').write_text(note,encoding='utf-8')


def defect_evidence():
    from competition_pipeline import DEFECTS
    a=pd.read_csv(ROOT/'data/1.csv').set_index('beer')
    b=pd.read_excel(ROOT/'data/8.xlsx',sheet_name='Calibrated chemical dataset').set_index('beer').reindex(a.index)
    meta=pd.read_excel(ROOT/'data/8.xlsx',sheet_name='Sensory thresholds');meta['field_name_R']=meta.field_name_R.astype('string').str.strip();meta=meta.set_index('field_name_R')
    rows=[]
    for f in DEFECTS:
        for beer in a.index:
            x,y=a.at[beer,f],b.at[beer,f]
            missx,missy=pd.isna(x),pd.isna(y)
            presence=bool(missx!=missy)
            conflict=bool(not missx and not missy and f in ['X2.3.butanedione','X.E..2.nonenal'] and not np.isclose(x,y,atol=1e-6,rtol=0))
            ambiguous=bool(not missy and y==0)
            tau=meta.at[f,'threshold_median']
            unknown=bool(missx or missy or ambiguous or conflict)
            status='native_missing' if missx else 'source_missing' if missy else 'numeric_source_conflict' if conflict else 'zero_semantics_unknown' if ambiguous else 'known_no_detected_source_conflict'
            rows.append(dict(beer=beer,field=f,a1_native=x,a8_concentration=y,unit=meta.at[f,'unit'],a1_missing=missx,a8_missing=missy,native_known=not missx,sources_presence_mismatch=presence,numeric_source_conflict=conflict,zero_semantics_unknown=ambiguous,threshold_median=tau,threshold_mean=meta.at[f,'threshold_mean'],threshold_unknown=pd.isna(tau),evidence_status=status,unknown=unknown,source_conflict=conflict,above_median_threshold=pd.NA if unknown or pd.isna(tau) else bool(y>tau),interpretation='感知阈值参考情景；冲突/零值/缺失不判为安全；非实际缺陷确认'))
    df=pd.DataFrame(rows);write(df,'q3_observed_defect_threshold_audit.csv')
    summary=[]
    for f,g in df.groupby('field',sort=False):
        summary.append(dict(field=f,n=len(g),native_known=int(g.native_known.sum()),native_missing=int(g.a1_missing.sum()),calibrated_missing=int(g.a8_missing.sum()),presence_mismatch=int(g.sources_presence_mismatch.sum()),numeric_source_conflict=int(g.numeric_source_conflict.sum()),zero_semantics_unknown=int(g.zero_semantics_unknown.sum()),threshold_unknown=int(g.threshold_unknown.sum()),interpretation='counts are overlapping evidence flags, not mutually exclusive risk classes'))
    write(pd.DataFrame(summary),'q3_defect_evidence_summary.csv')
    transferred=pd.read_csv(OUT/'q3_transferred_native_spectra.csv').set_index('source_beer')
    chosen=df.loc[df.beer.isin(transferred.index)].copy()
    chosen['source_beer']=chosen.beer
    chosen['before_native']=chosen.a1_native
    chosen['after_native']=[transferred.at[r.beer,r.field] for _,r in chosen.iterrows()]
    chosen['native_preserved']=np.isclose(chosen.before_native,chosen.after_native,atol=0,rtol=0,equal_nan=True)
    chosen['safety_claim']='measurement representation preserved only; unknown true concentration/threshold remains unknown'
    assert chosen.native_preserved.all()
    write(chosen,'q3_transfer_defect_evidence.csv')


def summary_consistency():
    fields=['esters_sum','aroma_hops_sum','acids_sum','sulfur_sum','kcalperc']
    meta=pd.read_excel(ROOT/'data/8.xlsx',sheet_name='Sensory thresholds');meta['field_name_R']=meta.field_name_R.astype('string').str.strip();meta=meta.set_index('field_name_R')
    source=pd.ExcelFile(FROZEN/'reference_supplement.xlsx')
    evidence=[]
    # Record source rows, not a guessed summation formula from names/correlations.
    for sheet in ['Supplementary Table S1','Sensory thresholds']:
        d=pd.read_excel(source,sheet_name=sheet)
        for i,row in d.iterrows():
            tokens=fields+['sum of acids','sum of esters','sum of hop','sum of sulfur','caloric value']
            if any(f.lower() in str(row.to_dict()).lower() for f in tokens):
                evidence.append({'sheet':sheet,'excel_row':i+2,'record':json.dumps(row.fillna('').to_dict(),ensure_ascii=False,default=str)})
    write(pd.DataFrame(evidence,columns=['sheet','excel_row','record']),'q3_summary_source_records.csv')
    registry=pd.read_csv(OUT/'calibrated_fields_transform_audit.csv') if (OUT/'calibrated_fields_transform_audit.csv').exists() else None
    fake=pd.read_csv(OUT/'q3_generated_native_spectra.csv')
    rows=[]
    for f in fields:
        # A workbook label "Total ..." does not establish constituents, coefficients,
        # assay, native-vs-concentration summation, or caloric conversion constants.
        unit=meta.at[f,'unit']
        transform='not_verified'
        if registry is not None:
            good=registry.loc[registry.field.eq(f)&registry.n.gt(1)&registry.matched.eq(registry.n)]
            if len(good)==1:transform='empirical_pairing_'+good.iloc[0].candidate_transform+'; not proof of constituent identity'
        rows.append(dict(field=f,unit=unit,formula_status='not_verified',constituent_list_status='not_verified',transform_evidence=transform,n_generated=len(fake),generated_missing=int(fake[f].isna().sum()),consistency_error=np.nan,recomputed=False,status='未验证',reason='No complete evidenced formula + constituents + units + transform chain in inspected sources; no numerical consistency error computed and no summary overwritten'))
    write(pd.DataFrame(rows),'q3_summary_consistency_audit.csv')


def guard_assessment():
    d=pd.read_csv(OUT/'q3_stale_guard_training_cv.csv')
    result={'folds':len(d),'model_mean_MAE':float(d.MAE.mean()),'baseline_mean_MAE':float(d.baseline_MAE.mean()),'paired_mean_MAE_difference':float((d.MAE-d.baseline_MAE).mean()),'model_better_folds':int(d.MAE.lt(d.baseline_MAE).sum()),'mean_R2':float(d.R2.mean()),'has_mean_MAE_advantage':bool(d.MAE.mean()<d.baseline_MAE.mean()),'selection_action':'none; auxiliary configuration and Q1 frozen selection unchanged','allowed_claim':'模型代理约束满足','disallowed_claim':'实际陈旧酒花缺陷不恶化','evidence_limit':'No predictive advantage shown over mean baseline in these fixed training folds; not an actual safety/quality measurement.'}
    (OUT/'q3_stale_guard_assessment.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')


def candidate_tables():
    """Problem-facing summaries, with no refitting or selection from held-out scores."""
    from competition_pipeline import load,DEFECTS
    x,y,styles,selection,tr,te=load()
    predictions=pd.read_csv(OUT/'q1_out_of_sample_predictions.csv')
    metrics=pd.read_csv(OUT/'q1_metrics.csv')
    groups={'Aroma':[t for t in y if t.startswith('A_')],'Flavour':[t for t in y if t.startswith('F_')],'All50':list(y)}
    rows=[]
    for split in ['training_oof','heldout_test']:
        for group,targets in groups.items():
            d=predictions.loc[predictions['split'].eq(split)&predictions.target.isin(targets)]
            a=d.actual_standardized.to_numpy();p=d.predicted_standardized.to_numpy();b=d.baseline_standardized.to_numpy()
            ag=d.actual_relative_grade.to_numpy();pg=d.predicted_relative_grade.to_numpy();bg=d.baseline_relative_grade.to_numpy()
            mae=np.abs(a-p).mean();base=np.abs(a-b).mean()
            rows.append(dict(split=split,group=group,n_targets=len(targets),n_beers=d.beer.nunique(),baseline_MAE=base,model_MAE=mae,improvement_pct=100*(1-mae/base),baseline_grade_MAE=np.abs(ag-bg).mean(),model_grade_MAE=np.abs(ag-pg).mean(),baseline_grade_accuracy=np.mean(ag==bg),model_grade_accuracy=np.mean(ag==pg)))
    write(pd.DataFrame(rows),'q1_group_metrics.csv')
    difficult=metrics.loc[metrics['split'].eq('heldout_test')].copy()
    difficult['mae_improvement_pct']=100*(1-difficult.MAE/difficult.baseline_MAE)
    difficult['comparison']=np.where(difficult.MAE>difficult.baseline_MAE+1e-10,'worse',np.where(difficult.MAE<difficult.baseline_MAE-1e-10,'better','equal'))
    write(difficult.sort_values('mae_improvement_pct'),'q1_dimension_comparison.csv')
    top=pd.read_csv(OUT/'q2_top3_all50_evidence.csv')
    nodes=pd.read_csv(OUT/'q2_graph_nodes.csv').set_index('field')
    edges=pd.read_csv(OUT/'q2_official_graph_edges.csv')
    edgekeys={frozenset([r.left,r.right]):r.rho for _,r in edges.iterrows()}
    context=[]
    for target,g in top.groupby('target',sort=False):
        names=g.field.tolist();present=[edgekeys[frozenset([a,b])] for i,a in enumerate(names) for b in names[i+1:] if frozenset([a,b]) in edgekeys]
        context.append(dict(target=target,top3_edges_at_abs_rho06=len(present),edge_rhos=json.dumps(present),role='auxiliary co-variation explanation; no network re-ranking or selection'))
    write(pd.DataFrame(context),'q2_network_explanation.csv')
    classes=top.groupby(['target','candidate_class']).size().unstack(fill_value=0).reset_index()
    write(classes,'q2_stability_by_target.csv')
    field_info=pd.read_csv(OUT/'chemical_name_mapping.csv').set_index('field_name_R')
    table=pd.read_excel(FROZEN/'reference_supplement.xlsx',sheet_name='Supplementary Table S1').rename(columns=lambda c:str(c).strip())
    names=table['Compound'].astype('string').str.strip().str.replace('Î±','α',regex=False).str.replace('Î²','β',regex=False)
    unique=table.loc[~names.duplicated(keep=False)].copy();unique.index=names.loc[unique.index]
    top['documented_compound_class']=[unique.at[field_info.at[f,'official_name'],'Compound class'] if field_info.at[f,'official_name'] in unique.index else pd.NA for f in top.field]
    write(top,'q2_top3_all50_evidence.csv')
    transfers=pd.read_csv(OUT/'q3_transferred_native_spectra.csv').set_index('source_beer')
    registry=pd.read_csv(OUT/'calibrated_fields_transform_audit.csv')
    native=[]
    for beer in transfers.index:
        for f in x.columns:
            before=x.at[beer,f];after=transfers.at[beer,f]
            r=registry.loc[registry.field.eq(f)&registry.n.gt(0)&registry.matched.eq(registry.n)]
            relation=r.iloc[0].candidate_transform if len(r)==1 else 'unverified_or_conflicted'
            native.append(dict(source_beer=beer,field=f,chemical_name=field_info.at[f,'official_name'],before_native=before,after_native=after,delta_native=after-before,delta_train_SD=(after-before)/x.loc[tr,f].std(ddof=0) if x.loc[tr,f].std(ddof=0)>0 else np.nan,unit_if_concentration=field_info.at[f,'unit'],native_transform_evidence=relation,is_defect=f in DEFECTS,interpretation='native coordinates; do not apply concentration units to log coordinates'))
    write(pd.DataFrame(native),'q3_transfer_compound_changes.csv')
    selected=['ethanol..v.v.','ethyl_acetate','ethyl_hexanoate','isoamyl_alcohol','bitterness','lactic_acid.mg.L.','acetic_acid.g.L.','pH']+DEFECTS
    d=pd.DataFrame(native)
    write(d.loc[d.field.isin(selected)].groupby(['field','chemical_name'],sort=False).agg(before_median=('before_native','median'),after_median=('after_native','median'),paired_delta_median=('delta_native','median'),known_pairs=('delta_native','count')).reset_index(),'q3_transfer_compound_summary.csv')
    sensory=pd.read_csv(OUT/'q3_transfer_sensory.csv')
    write(sensory.groupby('target',sort=False).agg(before_mean=('before','mean'),after_mean=('after','mean'),paired_mean_delta=('delta','mean')).reset_index(),'q3_transfer_sensory_summary.csv')
    design={'version':'v3','target_groups':groups,'Q1_selection':'frozen training-CV choices; no new tuning','network_role':'auxiliary visualization and co-variation context only; not an input to Top3 ranking','generator':'same-style anchor + one of five nearest same-style donors + Beta(1,4) interpolation; missing propagates','transfer_evaluation':'distance reduction is the optimization objective, not independent evidence of genuine style conversion','native_change_fields_predeclared':selected,'style_classifier':'existing fixed training-only diagnostic; not used to choose transfer candidates'}
    (OUT/'analysis_design.json').write_text(json.dumps(design,ensure_ascii=False,indent=2),encoding='utf-8')


def run_audits():
    from build_competition_deliverables import reference_audit
    reference_audit()
    scoring_sources();defect_evidence();summary_consistency();guard_assessment()
    diagnostics=pd.read_csv(OUT/'q3_generation_heldout_validation.csv')
    write(diagnostics,'q3_style_diversity_coverage.csv')
    candidate_tables()
    print('Review audits completed: score sources, four-defect flags, summary evidence, fixed guard CV, per-style coverage',flush=True)

if __name__=='__main__':run_audits()
