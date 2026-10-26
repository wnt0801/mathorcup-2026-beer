from pathlib import Path
import nbformat as nbf
from nbclient import NotebookClient

root = Path(__file__).resolve().parent
target = root / 'notebooks/05_defect_concentration_audit.ipynb'
if target.exists():
    raise FileExistsError(target)
nb = nbf.v4.new_notebook()
cells = []
def md(s): cells.append(nbf.v4.new_markdown_cell(s))
def code(s): cells.append(nbf.v4.new_code_cell(s))
md('''# 四种化合物浓度核验（独立审计）
范围：DMS、双乙酰、4VG、反式-2-壬烯醛；重点 beer=73、156、34、151。
仅读取原始文件，不删异常值、不填补、不替换来源、不训练或选择模型、不生成风格。
附件1是混合尺度，不能整表视为浓度。相对0–6级是训练集分位数分级，真实原始评分还原尚未解决。
本笔记本保存实际运行输出；所有数值判断均来自本地附件，不补造外部阈值。''')
code('''from pathlib import Path
import hashlib, json, sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from IPython.display import display
ROOT = Path.cwd() if (Path.cwd() / 'data').is_dir() else Path.cwd().parent
DATA = ROOT / 'data'
pd.set_option('display.max_columns', 30)
pd.set_option('display.max_rows', 100)
pd.set_option('display.precision', 12)
existing = sorted(p for folder in ['data','notebooks','outputs'] for p in (ROOT/folder).rglob('*') if p.is_file() and p.name != '05_defect_concentration_audit.ipynb')
before = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in existing}
display(pd.DataFrame([{'path': k, 'bytes': (ROOT/k).stat().st_size, 'sha256': v} for k,v in before.items()]))
print('Python:',sys.version)
print('数据说明：\\n'+(DATA/'数据说明.txt').read_text(encoding='utf-8'))
print('赛题文件：', [str(p.relative_to(ROOT)) for p in ROOT.rglob('*') if p.is_file() and '.venv' not in p.parts and (p.suffix.lower() in ['.pdf','.docx'] or '赛题' in p.name)])
''')
md('''## 已有材料与证据边界
当前工作区没有找到独立赛题正文（PDF、DOCX或文件名含“赛题”的文件），只能核对数据说明、README、已有 Notebook 和附件，不能声称已核验赛题原文。
01_data_check 第48–52单元已做初步变换、差异及阈值比检查；其中最后一格使用19、0.09、0.20作为阈值。下文核对它们的来源，不沿用其风险分类。
02_prediction、04_style_generation 当前为空白占位；03_compound_analysis 为已有关联分析。保留全部原文件。''')
code('''for p in sorted((ROOT/'notebooks').glob('*.ipynb')):
    if p.name == '05_defect_concentration_audit.ipynb': continue
    old = json.loads(p.read_text(encoding='utf-8'))
    print(p.name, 'cells=',len(old['cells']), 'saved_outputs=',sum(len(c.get('outputs',[])) for c in old['cells']))
    hits = [(i+1, ''.join(c['source'])) for i,c in enumerate(old['cells']) if any(t in ''.join(c['source']).lower() for t in ['threshold','nonenal','butanedione','check_features'])]
    for i,s in hits: print('相关单元',i, s[:6000])
print('README:'); print((ROOT/'README.md').read_text(encoding='utf-8'))
display(pd.DataFrame([{'file': p.name, 'rows':len(pd.read_csv(p)), 'columns':len(pd.read_csv(p).columns)} for p in sorted((ROOT/'outputs').glob('*.csv'))]))
''')
code('''a = pd.read_csv(DATA/'1.csv')
b = pd.read_excel(DATA/'8.xlsx',sheet_name='Calibrated chemical dataset')
t = pd.read_excel(DATA/'8.xlsx',sheet_name='Sensory thresholds')
fields = ['DMS','X2.3.butanedione','X4VG','X.E..2.nonenal']
assert a.beer.notna().all() and b.beer.notna().all()
assert a.beer.is_unique and b.beer.is_unique
assert set(a.beer) == set(b.beer) and len(a)==len(b)==250
a=a.set_index('beer'); b=b.set_index('beer').reindex(a.index)
assert a.beer_id.equals(b.beer_id)
print('附件1/8形状',a.shape,b.shape,'；250个beer唯一且集合相同，beer_id逐项一致')
t['field_name_R'] = t['field_name_R'].astype('string').str.strip()
meta=t.loc[t.field_name_R.isin(fields)]
assert len(meta)==4 and meta.field_name_R.is_unique
display(meta)
print('附件3工作表：', pd.ExcelFile(DATA/'3.xlsx').sheet_names)
print('附件8工作表：', pd.ExcelFile(DATA/'8.xlsx').sheet_names)
''')
md('''## 单位、零值、缺失与候选变换
单位来自附件8 Sensory thresholds 的 unit 字段：DMS为ug/L，其余三种为mg/L。这是附件元数据声明，不是独立实验校准证明；附件1的log值本身不能直接与浓度阈值比较。
只在附件8严格正值且两侧非缺失的配对上比较原值、ln、log10、log1p，另行完整列出零值及缺失。绝对误差容限1e-6，rtol=0；“匹配”仅表示数值关系。''')
code('''stats=[]; transforms=[]; issues=[]
for f in fields:
    x,y=a[f],b[f]
    for src,s in [('附件1',x),('附件8',y)]:
        stats.append(dict(field=f,source=src,n=len(s),missing=int(s.isna().sum()),zero=int(s.eq(0).sum()),negative=int(s.lt(0).sum()),infinite=int(np.isinf(s).sum()),minimum=s.min(),maximum=s.max()))
    mask=x.notna() & y.notna() & y.gt(0)
    for name,z in [('identity',y),('ln',np.log(y.where(y>0))),('log10',np.log10(y.where(y>0))),('log1p',np.log1p(y.where(y>=0)))]:
        delta=(x[mask]-z[mask]).abs()
        transforms.append(dict(field=f,transform=name,n=int(mask.sum()),matched=int(delta.le(1e-6).sum()),MAE=delta.mean(),max_abs_error=delta.max()))
    selected=x.isna()|y.isna()|x.eq(0)|y.eq(0)
    for beer in a.index[selected]: issues.append(dict(field=f,beer=beer,attachment1=x.loc[beer],attachment8=y.loc[beer],a1_missing=pd.isna(x.loc[beer]),a8_missing=pd.isna(y.loc[beer])))
display(pd.DataFrame(stats)); display(pd.DataFrame(transforms))
issues=pd.DataFrame(issues)
display(issues)
''')
code('''zero_evidence=[]
for f in ['DMS','X4VG']:
    x,y=a[f],b[f]
    zero=y.eq(0); positive=y.gt(0)
    floor=x[zero].unique()
    recovered=np.power(10.0,floor)
    minpos=y[positive].min()
    zero_evidence.append(dict(field=f,n_zero=int(zero.sum()),a1_zero_codes=floor.tolist(),inverse_log10=recovered.tolist(),a8_min_positive=minpos,half_min_positive=minpos/2,positive_shares_zero_code=int(x[positive].isin(floor).sum())))
display(pd.DataFrame(zero_evidence))
''')
md('''DMS的249个正值、4VG的82个正值均符合log10（最大误差约5e-10）。附件8零值分别1个、168个；附件1对应固定有限值，逆log10约等于该列最小正浓度的一半。这支持“零值经半最小正值替代后取log10”的数据关系，但不能证明生成脚本或零值的实验含义；不能据此认定为真实绝对零、低于检出限或缺失。原数据保持不变。''')
md('''## 原值对应字段的全部差异（含缺失状态差异）''')
code('''differences=[]
for f in ['X2.3.butanedione','X.E..2.nonenal']:
    same=np.isclose(a[f],b[f],atol=1e-6,rtol=0,equal_nan=True)
    for beer in a.index[~same]:
        x,y=a.at[beer,f],b.at[beer,f]
        differences.append(dict(field=f,beer=beer,attachment1=x,attachment8=y,delta_a1_minus_a8=x-y,kind='missing_state' if pd.isna(x) or pd.isna(y) else 'numeric_difference'))
diff=pd.DataFrame(differences)
display(diff)
focus = pd.concat([a.loc[[73,156,34,151],fields].assign(source='附件1'),b.loc[[73,156,34,151],fields].assign(source='附件8')]).reset_index()
display(focus)
display(b[['beer_id','tasting_category_fine','X2.3.butanedione']].nlargest(10,'X2.3.butanedione'))
print('beer=73 / 次高双乙酰：',b.at[73,'X2.3.butanedione']/b['X2.3.butanedione'].nlargest(2).iloc[1])
''')
md('''## 回到Excel原始单元格核对
核对源单元格地址、值、类型、公式/批注/格式，排除按行号读取和明显显示格式误解。没有公式或批注不代表数值已经通过实验质量控制。''')
code('''wb=load_workbook(DATA/'8.xlsx',data_only=False,read_only=False)
ws=wb['Calibrated chemical dataset']
headers={c.value:c.column for c in ws[1]}
rows={ws.cell(r,headers['beer']).value:r for r in range(2,ws.max_row+1)}
records=[]
for beer in [73,156,34,151,158,130,177]:
    for f in fields:
        c=ws.cell(rows[beer],headers[f])
        records.append(dict(beer=beer,field=f,cell=c.coordinate,value=c.value,data_type=c.data_type,number_format=c.number_format,comment=c.comment.text if c.comment else None))
display(pd.DataFrame(records))
wb.close()
''')
md('''## 可确定的结论与待决定事项

- **单位（附件声明）**：DMS ug/L；双乙酰、4VG、反式-2-壬烯醛 mg/L。DMS换算mg/L需除1000，但本次不改变任何值。
- **双乙酰**：241个共同非缺失值中238个按原值匹配，3个不匹配：beer156约为0.434130 / 0.123114；beer34约为0.289622 / 0.530466；beer151约为0.139752 / 0.043173（附件1 / 附件8）（精确值以上方运行输出为准）。附件1有9个缺失，附件8这9个位置全为0，不能自动当作已测得0，也不能直接判断哪份正确。
- **beer73极端值**：两份附件均为约121.731018 mg/L，附件8源格为数字而非公式。它是全列最高值，但另有数个高浓度样本；附件间一致不构成测量真实性证明。没有证据支持擅自除1000、删行或缩尾。
- **反式-2-壬烯醛**：249个共同有效值中247个按原值匹配。beer158、130数值不同；beer177附件1有值、附件8缺失。两份均无0，附件1无缺失，附件8有1个缺失。不能静默合并。
- **阈值溯源**：已有代码的19、0.09、0.20恰为附件8三项threshold_median；对应mean为17.25、0.29925、0.206333…。本次只核对来源，不把感知阈值直接定为缺陷判定阈值。反式-2-壬烯醛的mean、median均缺失，不补造。
- **下一步需决定**：差异样本采用哪一来源或维持冲突标记；9个“缺失/0”及DMS、4VG零值的实验语义；beer73及其他高值是否需原研究勘误或检测记录确认；缺陷定义、风格适用性及阈值依据。未解决前不做第三问风格生成模型。

以上小数以代码显示的源值为准；推断与已证实的数值关系分开。赛题正文未在当前工作区找到，需后续补齐才能核查题目对单位或阈值的额外限定。''')
code('''after={k:hashlib.sha256((ROOT/k).read_bytes()).hexdigest() for k in before}
assert before == after, '既有文件发生改变'
assert ((diff.field=='X2.3.butanedione') & (diff.kind=='numeric_difference')).sum()==3
assert set(diff.loc[(diff.field=='X2.3.butanedione') & (diff.kind=='numeric_difference'),'beer']) == {156,34,151}
print(f'核验完成：{len(before)}个既有数据、Notebook及输出文件的SHA256均未改变。未填补、未删行、未训练模型。')
''')
nb.cells=cells
nb.metadata['kernelspec']={'display_name':'Python 3','language':'python','name':'python3'}
nb.metadata['language_info']={'name':'python','version':'3.12'}
NotebookClient(nb,timeout=180,kernel_name='python3',resources={'metadata':{'path':str(root)}}).execute()
nbf.write(nb,target)
print(target)
